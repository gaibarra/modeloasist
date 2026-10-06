"""Reconcile semester workbooks by exact normalized name and campus, never source ID.

Run from backend: python -m app.services.semester_file_import --help
"""
import argparse
from collections import defaultdict
from datetime import datetime, time, timezone
from difflib import get_close_matches
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import select, text
from app.db.session import SessionLocal
import app.models.employee_credential  # register relationship before ORM configuration
from app.models import Employee, EmployeeDepartment, Department, StaffSemesterSchedule, StaffSemesterScheduleInterval, StaffScheduleFileImport

CAMPUS = {"CME": "Mérida", "CVA": "Valladolid", "CCH": "Chetumal"}
DAYS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def normalized(value):
    return " ".join("".join(c for c in unicodedata.normalize("NFKD", str(value).upper()) if not unicodedata.combining(c)).split())


def parse_time(value):
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError("Horas con segundos no soportadas")
        return value.hour * 60 + value.minute
    match = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?:\s*hrs)?\s*", str(value), re.I)
    if not match or int(match[1]) > 23 or int(match[2]) > 59:
        raise ValueError(f"Hora inválida: {value}")
    return int(match[1]) * 60 + int(match[2])


def merge_blocks(blocks):
    merged = []
    for start, end in sorted(set(blocks)):
        if start >= end:
            raise ValueError("Bloque sin duración o que cruza medianoche")
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return merged


def hhmm(minutes):
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def read_workbook(path):
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    groups = {}
    for sheet, campus in CAMPUS.items():
        ws = book[sheet]
        rows = ws.iter_rows(values_only=True)
        if tuple(next(rows))[:10] != ("empleadoId", "empleadoNombre", "escClave", "progClave", "matClave", "matNombre", "gpoClave", "diaLetra", "Inicio", "Fin"):
            raise ValueError(f"Encabezado inesperado en {sheet}")
        for number, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            key = (sheet, str(row[0]))
            person = groups.setdefault(key, {"sheet": sheet, "campus": campus, "source_id": str(row[0]), "name": str(row[1] or ""), "rows": [], "blocks": defaultdict(list), "errors": []})
            person["rows"].append(number)
            try:
                if not row[1] or normalized(row[1]) != normalized(person["name"]):
                    raise ValueError("ID de origen con nombres inconsistentes")
                weekday = [normalized(d) for d in DAYS].index(normalized(row[7]))
                start, end = parse_time(row[8]), parse_time(row[9])
                if start >= end:
                    raise ValueError("Intervalo inválido")
                person["blocks"][weekday].append((start, end))
            except ValueError as exc:
                person["errors"].append(f"Fila {number}: {exc}")
    book.close()
    result = []
    for person in groups.values():
        person["intervals"] = [{"weekday": day, "start": hhmm(start), "end": hhmm(end)} for day, blocks in sorted(person.pop("blocks").items()) for start, end in merge_blocks(blocks)]
        result.append(person)
    return result


def roster(db):
    employees = db.query(Employee).all()
    links = defaultdict(list)
    for link, department in db.query(EmployeeDepartment, Department).join(Department, Department.id == EmployeeDepartment.department_id):
        if department.active:
            links[link.employee_id].append(department)
    result = []
    for employee in employees:
        campuses = {normalized(d.campus) for d in links[employee.id] if d.campus}
        if employee.campus:
            campuses.add(normalized(employee.campus))
        if normalized(employee.campus or "") == "BAJAS":
            continue
        covered = [c for c in CAMPUS.values() if normalized(c) in campuses]
        if covered:
            result.append({"id": employee.id, "name": employee.nombre, "email": employee.email, "campuses": covered,
                           "departments": ", ".join(sorted({d.name for d in links[employee.id]})) or employee.departamento})
    return result


def reconcile(people, employees):
    matches, issues = [], []
    for person in people:
        candidates = [e for e in employees if person["campus"] in e["campuses"] and normalized(person["name"]) == normalized(e["name"])]
        if normalized(person["name"]) == "SIN MAESTRO":
            reason = "Marcador sin persona asignada"
        elif person["errors"]:
            reason = "; ".join(person["errors"])
        elif len(candidates) != 1:
            reason = "Sin coincidencia exacta de nombre y campus" if not candidates else "Nombre ambiguo"
        else:
            matches.append({**person, "employee_id": candidates[0]["id"], "employee_name": candidates[0]["name"]})
            continue
        names = {normalized(e["name"]): e["name"] for e in employees if person["campus"] in e["campuses"]}
        suggestions = [names[n] for n in get_close_matches(normalized(person["name"]), names, n=3, cutoff=0.8)]
        issues.append({**person, "reason": reason, "suggestions": suggestions})
    # One employee can have multiple external IDs / sheets: combine validated blocks.
    grouped = defaultdict(list)
    for match in matches:
        grouped[match["employee_id"]].append(match)
    updates = []
    for employee_id, sources in grouped.items():
        blocks = defaultdict(list)
        for source in sources:
            for interval in source["intervals"]:
                blocks[interval["weekday"]].append((parse_time(interval["start"]), parse_time(interval["end"])))
        # Do not partially confirm an employee if another source row is invalid.
        if any(normalized(i["name"]) == normalized(s["name"]) and i["campus"] == s["campus"] for i in issues for s in sources):
            issues.extend({**s, "reason": "Otra fuente de esta persona requiere revisión", "suggestions": []} for s in sources)
            continue
        updates.append({"employee_id": employee_id, "name": sources[0]["employee_name"], "sources": sources,
                        "intervals": [{"weekday": day, "start": hhmm(a), "end": hhmm(b)} for day, pairs in sorted(blocks.items()) for a, b in merge_blocks(pairs)]})
    return sorted(updates, key=lambda u: normalized(u["name"])), issues


def report(path, audit, employees, applied):
    confirmed = {u["employee_id"] for u in audit["updates"]}
    book = openpyxl.Workbook()
    overview = book.active
    overview.title = "Resumen"
    overview.append(["Segundo semestre 2026", "Agosto–Diciembre"])
    overview.append(["Estado", "Aplicado" if applied else "Vista previa; sin cambios"])
    overview.append(["Fuente", audit["source_name"]])
    overview.append(["SHA256", audit["sha256"]])
    overview.append(["Generado UTC", datetime.now(timezone.utc).isoformat()])
    overview.append(["Confirmados por archivo", len(confirmed)])
    overview.append(["Pendientes de confirmar", sum(e["id"] not in confirmed for e in employees)])
    overview.append(["Fuentes por revisar", len(audit["issues"])])
    overview.append(["Criterio", audit.get("matching_criterion", "Coincidencia exacta de nombre normalizado + campus; IDs externos solo informativos.")])
    overview.append(["Pendientes", "Conservan su horario previo; no equivalen a descanso ni a ausencia."])
    overview.append(["Solicitud", "Favor de confirmar horario agosto–diciembre 2026: bloques de entrada/salida por día, descansos y fecha de vigencia."])
    for campus in CAMPUS.values():
        sheet = book.create_sheet("Pendientes " + campus)
        sheet.append(["ID asistencia", "Colaborador", "Correo", "Departamento", "Campus", "Estado", "Horario confirmado recibido", "Observaciones"])
        for e in sorted(employees, key=lambda e: normalized(e["name"])):
            if e["id"] not in confirmed and campus in e["campuses"]:
                sheet.append([e["id"], e["name"], e["email"], e["departments"], campus, "Pendiente de confirmación", "", "Solicitar horario completo del semestre"])
    sheet = book.create_sheet("Confirmados")
    sheet.append(["ID asistencia", "Colaborador", "Campus fuente", "ID externo", *DAYS])
    for u in audit["updates"]:
        sheet.append([u["employee_id"], u["name"], ", ".join(sorted({s["campus"] for s in u["sources"]})), ", ".join(s["source_id"] for s in u["sources"]), *[" / ".join(f'{i["start"]}–{i["end"]}' for i in u["intervals"] if i["weekday"] == d) or "Sin turno" for d in range(7)]])
    sheet = book.create_sheet("Revisar coincidencias")
    sheet.append(["Pestaña", "ID externo", "Nombre archivo", "Filas", "Motivo", "Sugerencias (no aplicadas)"])
    for i in audit["issues"]:
        sheet.append([i["sheet"], i["source_id"], i["name"], ",".join(map(str, i["rows"])), i["reason"], " / ".join(i["suggestions"])])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="163D68")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(48, max(16, max(len(str(c.value or "")) for c in column) + 2))
            for cell in column:
                # Prevent spreadsheet formula interpretation of names / free text.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    book.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    sha = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    people = read_workbook(args.workbook)
    args.output.mkdir(parents=True, exist_ok=True)
    with SessionLocal() as db:
        employees = roster(db)
        updates, issues = reconcile(people, employees)
        audit = {"source_name": args.workbook.name, "sha256": sha, "year": 2026, "semester": 2, "updates": updates, "issues": issues}
        applied = False
        if args.apply:
            db.execute(text("SELECT pg_advisory_xact_lock(20260914)"))
            receipt = db.query(StaffScheduleFileImport).filter_by(source_sha256=sha, academic_year=2026, semester=2).first()
            if receipt:
                audit = receipt.audit
            else:
                for u in updates:
                    profile = db.query(StaffSemesterSchedule).filter_by(employee_id=u["employee_id"], academic_year=2026, semester=2).with_for_update().first()
                    u["previous"] = None if profile is None else {"id": profile.id, "updated_by_staff_user_id": profile.updated_by_staff_user_id, "updated_at": str(profile.updated_at), "intervals": [{"weekday": i.weekday, "start": i.start.isoformat(), "end": i.end.isoformat()} for i in profile.intervals]}
                # Recovery snapshot is durable before any schedule is replaced.
                backup = args.output / f"respaldo-{sha[:12]}-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%f')}.json"
                with backup.open("x") as handle:
                    json.dump(audit, handle, ensure_ascii=False, indent=2)
                for u in updates:
                    profile = db.query(StaffSemesterSchedule).filter_by(employee_id=u["employee_id"], academic_year=2026, semester=2).first()
                    if profile is None:
                        profile = StaffSemesterSchedule(employee_id=u["employee_id"], academic_year=2026, semester=2)
                        db.add(profile)
                    profile.intervals = [StaffSemesterScheduleInterval(weekday=i["weekday"], start=time.fromisoformat(i["start"]), end=time.fromisoformat(i["end"])) for i in u["intervals"]]
                    profile.updated_at = datetime.now(timezone.utc)
                    profile.updated_by_staff_user_id = None  # file import, not an impersonated staff action
                db.add(StaffScheduleFileImport(source_sha256=sha, source_name=args.workbook.name, academic_year=2026, semester=2, audit=audit))
                db.flush()
                for u in updates:
                    profile = db.query(StaffSemesterSchedule).filter_by(employee_id=u["employee_id"], academic_year=2026, semester=2).one()
                    actual = sorted((i.weekday, i.start.strftime("%H:%M"), i.end.strftime("%H:%M")) for i in profile.intervals)
                    expected = sorted((i["weekday"], i["start"], i["end"]) for i in u["intervals"])
                    if actual != expected:
                        raise RuntimeError(f'Verificación fallida: {u["employee_id"]}')
                db.commit()
            applied = True
        report(args.output / "Seguimiento-horarios-2026-S2.xlsx", audit, employees, applied)
        confirmed = {u["employee_id"] for u in audit["updates"]}
        print(json.dumps({"applied": applied, "confirmed": len(confirmed), "pending": sum(e["id"] not in confirmed for e in employees), "review": len(audit["issues"]), "by_campus": {campus: {"confirmed": sum(e["id"] in confirmed and campus in e["campuses"] for e in employees), "pending": sum(e["id"] not in confirmed and campus in e["campuses"] for e in employees)} for campus in CAMPUS.values()}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
