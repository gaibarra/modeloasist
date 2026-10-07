"""Operator-only preview/apply/rollback. Preview never writes to the database.

python -m app.services.labor_contract_import --help
Sensitive RFC/CURP/IMSS/hire-date columns are deliberately never retained.
"""
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
import hashlib
import json
import os
from pathlib import Path
import subprocess
import unicodedata

from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import selectinload
from app.models import Employee, EmployeeDepartment, Department, StaffUser, AttendanceImportBatch, StaffSemesterSchedule, LaborContract, LaborContractImport
import app.models.employee_credential  # register relationships
from app.services.labor_contracts import available, semester_minutes, comparison, LABELS

START, END = date(2026, 8, 1), date(2026, 12, 31)
CAMPUS = {"escuela": "Montejo", "universidad": "Mérida", "valladolid": "Valladolid", "chetumal": "Chetumal"}
FIELDS = ("nombre", "email", "campus", "departamento")


def normalize(value):
    return " ".join("".join(c for c in unicodedata.normalize("NFD", str(value or "").upper()) if not unicodedata.combining(c)).split())


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def read_source(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = []
    try:
        if {s.title.strip().lower() for s in workbook} != set(CAMPUS):
            raise ValueError("Las cuatro pestañas no coinciden con el formato aprobado")
        for sheet in workbook:
            title = sheet.title.strip().lower()
            values = sheet.iter_rows(values_only=True)
            header = next(values)
            if len(header) < 10 or normalize(header[1]) != "NOMBRE COMPLETO" or normalize(header[9]) != "HORAS CONTRATADAS":
                raise ValueError("Encabezados inesperados")
            for number, data in enumerate(values, 2):
                if not any(v is not None for v in data):
                    continue
                minutes, errors = None, []
                if not data[1] or data[0] is None:
                    errors.append("Falta clave o nombre")
                if data[9] is not None:
                    try:
                        value = Decimal(str(data[9])) * 60
                        if not value.is_finite() or value < 0 or value > 10080 or value != value.to_integral_value():
                            raise ValueError()
                        minutes = int(value)
                    except (ValueError, InvalidOperation):
                        errors.append("Horas inválidas: se requieren horas semanales con precisión de minutos")
                rows.append(dict(key=f"{title}:{number}", sheet=title, row=number,
                                 source_key=str(data[0] or ""), source_name=str(data[1] or "").strip(),
                                 email=str(data[4] or "").strip().lower(), department_label=str(data[5] or ""),
                                 department_code=str(data[6] or ""), campus=CAMPUS[title],
                                 weekly_minutes=minutes, errors=errors))
    finally:
        workbook.close()
    for field in ("source_key", "source_name", "email"):
        counts = Counter(normalize(r[field]) for r in rows if r[field])
        for row in rows:
            if row[field] and counts[normalize(row[field])] > 1:
                row["errors"].append(f"Duplicado en archivo: {field}")
    return rows


def profile(db, employee):
    return {**{field: getattr(employee, field) for field in FIELDS},
            "department_ids": sorted(db.scalars(select(EmployeeDepartment.department_id).where(EmployeeDepartment.employee_id == employee.id)).all())}


def state_digest(db):
    links = defaultdict(list)
    for eid, did in db.execute(select(EmployeeDepartment.employee_id, EmployeeDepartment.department_id)):
        links[eid].append(did)
    employees = [{"id": e.id, **{f: getattr(e, f) for f in FIELDS}, "department_ids": sorted(links[e.id])} for e in db.scalars(select(Employee).order_by(Employee.id))]
    schedules = [{"id": s.id, "employee_id": s.employee_id, "year": s.academic_year, "semester": s.semester,
                  "blocks": sorted((b.weekday, str(b.start), str(b.end)) for b in s.intervals)}
                 for s in db.scalars(select(StaffSemesterSchedule).options(selectinload(StaffSemesterSchedule.intervals)).order_by(StaffSemesterSchedule.id))]
    contracts = []
    if available(db):
        contracts = [(c.id, c.employee_id, str(c.valid_from), str(c.valid_until), c.weekly_minutes, c.active)
                     for c in db.scalars(select(LaborContract).order_by(LaborContract.id))]
    return digest([employees, schedules, contracts])


def preview(db, path, department_map=None, confirmed_matches=None):
    department_map, confirmed_matches = department_map or {}, confirmed_matches or {}
    rows = read_source(path)
    people = list(db.scalars(select(Employee).order_by(Employee.id)))
    by_id = {e.id: e for e in people}
    names, emails = defaultdict(set), defaultdict(set)
    for employee in people:
        names[normalize(employee.nombre)].add(employee.id)
        if employee.email:
            emails[normalize(employee.email)].add(employee.id)
    provisional_names = defaultdict(set)
    for batch in db.scalars(select(AttendanceImportBatch)):
        for item in batch.auto_created_employees or []:
            provisional_names[item["employee_id"]].add(normalize(item["nombre"]))
    departments = {d.id: d for d in db.scalars(select(Department))}
    schedules = {s.employee_id: s for s in db.scalars(select(StaffSemesterSchedule).where(StaffSemesterSchedule.academic_year == 2026, StaffSemesterSchedule.semester == 2))}
    for row in rows:
        a, b = names[normalize(row["source_name"])], emails[normalize(row["email"])] if row["email"] else set()
        candidates = a | b
        row.update(employee_id=None, application_name=None, changes={}, add_department_id=None, warnings=[], candidate_ids=sorted(candidates), candidates=[])
        if row["key"] in confirmed_matches:
            eid = confirmed_matches[row["key"]]
            if eid not in by_id:
                row["errors"].append("Confirmación refiere a una persona inexistente")
            elif b and b != {eid}:
                row["errors"].append("El correo pertenece a otra cuenta: resolver antes de confirmar")
            else:
                row["employee_id"] = eid
                row["match"] = "Confirmación explícita"
        elif len(candidates) == 1:
            row["employee_id"] = next(iter(candidates))
            row["match"] = "Nombre/correo único"
        else:
            row["match"] = "Ambiguo" if candidates else "Sin coincidencia exacta"
            if not candidates:
                scored = sorted(((SequenceMatcher(None, normalize(row["source_name"]), normalize(e.nombre)).ratio(), e.id) for e in people), reverse=True)
                candidates = {eid for score, eid in scored[:3] if score >= .78}
            row["candidate_ids"] = sorted(candidates)
            row["candidates"] = [{"employee_id": eid, "name": by_id[eid].nombre} for eid in sorted(candidates)]
        eid = row["employee_id"]
        if eid is None:
            row["status"] = "pending_identity"
            continue
        employee = by_id[eid]
        row["application_name"] = employee.nombre
        before = profile(db, employee)
        row["before"] = before
        placeholder = not employee.email or employee.email.lower().endswith("@pendiente.local")
        if row["email"] and normalize(row["email"]) != normalize(employee.email):
            if placeholder and not emails[normalize(row["email"])]:
                if "@" not in row["email"] or " " in row["email"]:
                    row["warnings"].append("Correo de origen inválido; no se cambiará")
                else:
                    row["changes"]["email"] = row["email"]
                    row["warnings"].append("Cambiará el correo de acceso; se conserva la contraseña")
            else:
                row["warnings"].append("Correo confirmado distinto: se conserva")
        if normalize(row["source_name"]) != normalize(employee.nombre):
            if not employee.nombre or (placeholder and normalize(employee.nombre) in provisional_names[eid]):
                row["changes"]["nombre"] = row["source_name"]
            else:
                row["warnings"].append("Nombre confirmado distinto: se conserva")
        linked_campuses = {departments[d].campus for d in before["department_ids"] if departments[d].campus}
        if not employee.campus and (not linked_campuses or linked_campuses == {row["campus"]}):
            row["changes"]["campus"] = row["campus"]
        elif employee.campus and employee.campus != row["campus"]:
            row["warnings"].append("Campus distinto: se conserva el confirmado")
        map_key = f"{row['sheet']}:{row['department_code']}"
        mapped = department_map.get(map_key)
        if mapped is not None:
            department = departments.get(mapped)
            if not department or department.campus != row["campus"]:
                row["errors"].append("Mapeo de departamento inválido o de otro campus")
            elif not before["department_ids"]:
                row["add_department_id"] = mapped
                if not employee.departamento:
                    row["changes"]["departamento"] = department.name
        elif not before["department_ids"]:
            row["warnings"].append("Departamento pendiente de mapeo explícito")
        row["scheduled_minutes"] = semester_minutes(schedules.get(eid))
        row["status"] = comparison(row["weekly_minutes"], row["scheduled_minutes"])
    targets = Counter(r["employee_id"] for r in rows if r["employee_id"] is not None)
    for row in rows:
        if row["employee_id"] is not None and targets[row["employee_id"]] > 1:
            row["errors"].append("Dos filas apuntan al mismo colaborador")
        row["eligible"] = row["employee_id"] is not None and not row["errors"]
        row["status_label"] = LABELS[row["status"]]
    seen = {r["employee_id"] for r in rows if r["employee_id"]}
    mapping_review = {}
    for row in rows:
        key = f"{row['sheet']}:{row['department_code']}"
        entry = mapping_review.setdefault(key, {"source_department": row["department_label"], "campus": row["campus"], "candidate_departments": {}})
        for did in row.get("before", {}).get("department_ids", []):
            entry["candidate_departments"][str(did)] = departments[did].name
    return {"version": 1, "source_name": Path(path).name, "source_sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
            "valid_from": str(START), "valid_until": str(END), "state_sha256": state_digest(db),
            "department_map": department_map, "confirmed_matches": confirmed_matches, "rows": rows,
            "department_mapping_review": mapping_review,
            "catalog_unmatched": [{"employee_id": e.id, "name": e.nombre, "campus": e.campus} for e in people if e.id not in seen],
            "summary": {"source_rows": len(rows), "eligible": sum(r["eligible"] for r in rows),
                        "pending_identity": sum(r["employee_id"] is None for r in rows),
                        "invalid": sum(bool(r["errors"]) for r in rows),
                        "missing_hours": sum(r["weekly_minutes"] is None for r in rows),
                        "profile_updates": sum(bool(r["changes"] or r["add_department_id"]) for r in rows),
                        "statuses": dict(Counter(r["status_label"] for r in rows)),
                        "catalog_not_matched": len(by_id.keys() - seen)}}


def require_operator(db, staff_id):
    staff = db.get(StaffUser, staff_id)
    if not staff or not staff.is_active or not staff.is_superadmin:
        raise ValueError("La importación global requiere un staff superadministrador activo")


def active_ids(db, employee_id):
    return sorted(db.scalars(select(LaborContract.id).where(LaborContract.employee_id == employee_id, LaborContract.active.is_(True))).all())


def contract_state(db, employee_id):
    return [{"id": c.id, "start": str(c.valid_from), "end": str(c.valid_until), "minutes": c.weekly_minutes, "active": c.active}
            for c in db.scalars(select(LaborContract).where(LaborContract.employee_id == employee_id).order_by(LaborContract.id))]


def lock_catalog(db):
    if db.bind.dialect.name == "postgresql":
        # Serialize imports and prevent concurrent profile/schedule edits during validation.
        db.execute(text("SET LOCAL lock_timeout = '5s'"))
        db.execute(text("LOCK TABLE employees, employee_departments, staff_semester_schedules, staff_semester_schedule_intervals, labor_contracts, labor_contract_imports IN SHARE ROW EXCLUSIVE MODE"))


def apply_preview(db, path, saved, approved_keys, staff_id):
    require_operator(db, staff_id)
    if not available(db):
        raise ValueError("Aplica primero la migración de contratos con respaldo")
    if not approved_keys or len(approved_keys) != len(set(approved_keys)):
        raise ValueError("Se requiere una lista explícita de filas aprobadas, sin duplicados")
    request_hash = digest([saved["source_sha256"], sorted(approved_keys), saved["confirmed_matches"], saved["department_map"]])
    lock_catalog(db)
    db.expire_all()
    previous = db.scalar(select(LaborContractImport).where(LaborContractImport.request_sha256 == request_hash))
    if previous:
        if previous.reverted_at:
            raise ValueError("La operación ya fue revertida; requiere un nuevo proceso de revisión")
        return previous.id
    fresh = preview(db, path, saved["department_map"], saved["confirmed_matches"])
    if digest(fresh) != digest(saved):
        raise ValueError("La fuente, vista previa o el catálogo cambiaron. Genera y confirma otra vista previa")
    indexed = {r["key"]: r for r in fresh["rows"]}
    if any(k not in indexed or not indexed[k]["eligible"] for k in approved_keys):
        raise ValueError("La selección contiene filas pendientes o inválidas")
    # A source row may only be applied once, even through differently grouped requests.
    applied = {r["key"] for receipt in db.scalars(select(LaborContractImport).where(LaborContractImport.source_sha256 == fresh["source_sha256"], LaborContractImport.reverted_at.is_(None))) for r in receipt.audit.get("changes", [])}
    if applied.intersection(approved_keys):
        raise ValueError("La selección incluye filas de este archivo ya aplicadas")
    receipt = LaborContractImport(source_sha256=fresh["source_sha256"], request_sha256=request_hash,
                                  source_name=fresh["source_name"], staff_user_id=staff_id, audit={})
    db.add(receipt)
    db.flush()
    changes = []
    for key in approved_keys:
        row = indexed[key]
        employee = db.get(Employee, row["employee_id"])
        old_ids = active_ids(db, employee.id)
        replaced = []
        for field, value in row["changes"].items():
            setattr(employee, field, value)
        if row["add_department_id"]:
            db.add(EmployeeDepartment(employee_id=employee.id, department_id=row["add_department_id"], is_primary=True))
        new_id = None
        if row["weekly_minutes"] is not None:
            overlaps = list(db.scalars(select(LaborContract).where(LaborContract.employee_id == employee.id, LaborContract.active.is_(True), LaborContract.valid_from <= END, LaborContract.valid_until >= START)))
            if any(c.valid_from != START or c.valid_until != END for c in overlaps):
                raise ValueError("Contrato con vigencia solapada; requiere revisión manual")
            for contract in overlaps:
                contract.active = False
                replaced.append(contract.id)
            contract = LaborContract(employee_id=employee.id, valid_from=START, valid_until=END,
                                     weekly_minutes=row["weekly_minutes"], source_sheet=row["sheet"], source_row=row["row"], import_id=receipt.id)
            db.add(contract)
            db.flush()
            new_id = contract.id
        db.flush()
        changes.append({"key": key, "employee_id": employee.id, "before": row["before"], "after": profile(db, employee),
                        "old_contract_ids": old_ids, "after_contract_ids": active_ids(db, employee.id),
                        "after_contract_state": contract_state(db, employee.id),
                        "replaced_contract_ids": replaced, "new_contract_id": new_id})
    receipt.audit = {"source_rows": [{"key": r["key"], "employee_id": r["employee_id"], "candidate_ids": r["candidate_ids"],
                                      "source_key": r["source_key"], "department_code": r["department_code"]} for r in fresh["rows"]],
                     "changes": changes, "valid_from": str(START), "valid_until": str(END)}
    db.flush()
    return receipt.id


def rollback(db, import_id, staff_id):
    require_operator(db, staff_id)
    lock_catalog(db)
    db.expire_all()
    receipt = db.get(LaborContractImport, import_id)
    if not receipt or receipt.reverted_at:
        raise ValueError("Operación inexistente o ya revertida")
    for change in receipt.audit["changes"]:
        employee = db.get(Employee, change["employee_id"])
        if not employee or profile(db, employee) != change["after"] or contract_state(db, employee.id) != change["after_contract_state"]:
            raise ValueError("Hay cambios posteriores; no se permite sobrescribirlos")
    for change in receipt.audit["changes"]:
        employee = db.get(Employee, change["employee_id"])
        for field in FIELDS:
            setattr(employee, field, change["before"][field])
        added = set(change["after"]["department_ids"]) - set(change["before"]["department_ids"])
        for link in db.scalars(select(EmployeeDepartment).where(EmployeeDepartment.employee_id == employee.id, EmployeeDepartment.department_id.in_(added))):
            db.delete(link)
        if change["new_contract_id"]:
            db.get(LaborContract, change["new_contract_id"]).active = False
        for cid in change["replaced_contract_ids"]:
            db.get(LaborContract, cid).active = True
    receipt.reverted_at, receipt.reverted_by = datetime.now(timezone.utc), staff_id
    db.flush()


def save_private(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "w") as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2, default=str)


def export_review(path, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = "Vista previa"
    sheet.append(["Fila origen", "Clave origen (no ID)", "Nombre Excel", "ID aplicación", "Nombre aplicación", "Identificación", "Horas contrato", "Horas horario", "Estado", "Cambios propuestos", "Advertencias", "Errores", "Candidatos", "Apta para aprobar"])
    def safe(value):
        return "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value
    for row in data["rows"]:
        sheet.append([safe(v) for v in [row["key"], row["source_key"], row["source_name"], row["employee_id"], row["application_name"], row["match"],
                     row["weekly_minutes"] / 60 if row["weekly_minutes"] is not None else None,
                     row.get("scheduled_minutes") / 60 if row.get("scheduled_minutes") is not None else None,
                     row["status_label"], json.dumps(row["changes"], ensure_ascii=False), "; ".join(row["warnings"]), "; ".join(row["errors"]),
                     "; ".join(f"{c['employee_id']}: {c['name']}" for c in row["candidates"]), "Sí" if row["eligible"] else "No"]])
    absent = book.create_sheet("Catálogo sin coincidencia")
    absent.append(["ID", "Nombre aplicación", "Campus"])
    for row in data["catalog_unmatched"]:
        absent.append([safe(row[k]) for k in ("employee_id", "name", "campus")])
    mapping = book.create_sheet("Mapeo por confirmar")
    mapping.append(["Hoja:clave departamento", "Descripción origen", "Campus", "Departamentos de candidatos (no asignados automáticamente)"])
    for key, value in data["department_mapping_review"].items():
        mapping.append([safe(key), safe(value["source_department"]), value["campus"], json.dumps(value["candidate_departments"], ensure_ascii=False)])
    for ws in book:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365C")
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(60, max(18, max(len(str(c.value or "")) for c in column) + 2))
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "wb") as stream:
        book.save(stream)


def backup_database(engine, target):
    if engine.dialect.name != "postgresql":
        raise ValueError("La aplicación operativa requiere PostgreSQL")
    url = engine.url
    env = os.environ.copy()
    env.update(PGHOST=url.host or "localhost", PGPORT=str(url.port or 5432), PGUSER=url.username or "", PGPASSWORD=url.password or "", PGDATABASE=url.database or "")
    for key in ("sslmode", "sslrootcert", "sslcert", "sslkey", "options"):
        if key in url.query:
            env["PG" + key.upper()] = str(url.query[key])
    fd = os.open(target, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "wb") as stream:
        subprocess.run(["pg_dump", "--format=custom", "--no-owner", "--no-acl"], env=env, stdout=stream, check=True)
    # Debian's pg_wrapper selects the client using PGHOST/PGPORT. Use the same
    # environment for both tools, otherwise restore may pick an older client.
    subprocess.run(["pg_restore", "--list", str(target)], env=env, check=True, stdout=subprocess.DEVNULL)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["preview", "apply", "rollback"])
    parser.add_argument("--source", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--preview", type=Path)
    parser.add_argument("--mapping", type=Path, help='JSON {"department_map": {"sheet:code": id}, "confirmed_matches": {"sheet:row": employee_id}}')
    parser.add_argument("--approved", type=Path, help="JSON array con claves sheet:row explícitamente aprobadas")
    parser.add_argument("--staff-id", type=int)
    parser.add_argument("--import-id", type=int)
    parser.add_argument("--backup", type=Path)
    args = parser.parse_args()
    from app.db.session import SessionLocal
    with SessionLocal() as db:
        if args.action == "preview":
            if not args.source or not args.output:
                parser.error("preview requiere --source y --output")
            if db.bind.dialect.name == "postgresql":
                db.execute(text("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY"))
            config = json.loads(args.mapping.read_text()) if args.mapping else {}
            result = preview(db, args.source, config.get("department_map"), config.get("confirmed_matches"))
            save_private(args.output, result)
            export_review(args.output.with_suffix(".xlsx"), result)
            print(json.dumps(result["summary"], ensure_ascii=False, indent=2))
            return
        if not args.staff_id or not args.backup:
            parser.error("Las escrituras requieren --staff-id y --backup (archivo nuevo)")
        require_operator(db, args.staff_id)
        db.rollback()
        backup_database(db.bind, args.backup)
        with db.begin():
            if args.action == "apply":
                if not all([args.source, args.preview, args.approved]):
                    parser.error("apply requiere --source, --preview y --approved")
                result = apply_preview(db, args.source, json.loads(args.preview.read_text()), json.loads(args.approved.read_text()), args.staff_id)
            else:
                if not args.import_id:
                    parser.error("rollback requiere --import-id")
                rollback(db, args.import_id, args.staff_id)
                result = args.import_id
        print(f"Operación {result}: {args.action} completada. Respaldo: {args.backup}")


if __name__ == "__main__":
    main()
